import datetime
import io
import zipfile
from decimal import Decimal
from unittest.mock import DEFAULT, Mock, patch

import openpyxl
from django.apps import apps as django_apps
from django.conf import settings
from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.http import FileResponse
from django.test import RequestFactory, tag
from django.urls import reverse

import geno.tests.data as geno_testdata
from geno.models import (
    Document,
    DocumentType,
    Invoice,
    InvoiceCategory,
    Share,
    ShareType,
)
from geno.views import InvoiceBatchGenerateView, ShareStatementView

from .base import GenoAdminTestCase, fill_template_effect


def create_fake_pdf(
    _ref_number,
    _address,
    _context,
    output_filename,
    _render,
    _email_template,
    _email_subject,
    _dry_run,
):
    with open(f"/tmp/{output_filename}", "wb") as dummy_pdf:
        dummy_pdf.write(b"Dummy PDF")
    return DEFAULT


class GenoViewsTest(GenoAdminTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        geno_testdata.create_contracts(cls)

    def test_views_status200(self):
        logged_in = self.client.login(username="superuser", password="secret")
        self.assertTrue(logged_in)
        paths = [
            #'import/([a-z_-]+)/',
            "export/adit/",
            "share/overview/",
            # ADDED BELOW 'share/overview/plot/',
            "share/export/",
            "share/confirm/",
            reverse("geno:share-statement-form"),
            reverse("geno:share-statement", args=["current_year"]),
            reverse("geno:share-statement-for-address", args=["current_year", 1]),
            # TODO: do we need the share/mailing/ endpoint?
            # "share/mailing/",
            "share/interest/download/",
            "share/interest/transactions/",
            #'share/interest/adjust/',
            "share/duedate_reminder/",
            "member/overview/",
            # ADDED BELOW 'member/overview/plot/',
            "member/list/",
            "member/list_admin/",
            # TODO 'member/check_mailinglists/',
            "member/check_payments/",
            reverse("geno:mail-wizard-start"),
            # NEEDS SESSION reverse("geno:mail-wizard-select"),
            reverse("geno:mail-wizard-action"),
            "member/confirm/memberletter/",
            "address/export/",
            "maintenance/",
            # ADDED BELOW 'maintenance/check_portal_users/',
            "transaction/",
            reverse("geno:transaction-upload"),
            reverse("geno:transaction-process"),
            reverse("geno:transaction-testdata"),
            reverse("geno:transaction-invoice"),
            reverse("geno:invoice-manual"),
            reverse("geno:invoice-batch"),
            reverse("geno:invoice-batch-generate"),
            reverse("geno:invoice-download", args=["contract", 1]),
            reverse("geno:debtor-list"),
            # TODO: NEEDS DATA'debtor/detail/contract/1/',
            # TODO: OLD CODE'contract/create/',
            # TODO: OLD CODE'contract/create_letter/',
            reverse("geno:contract-check-forms"),
            reverse("geno:contract-report"),
            "rental/units/",
            "rental/units/mailbox/",
            "rental/units/protocol/",
            "odt2pdf/",
            #'webstamp/',
            #'oauth_client/',
            #'oauth_client/login'
            #'oauth_client/callback/',
            #'oauth_client/test/',
            "preview/",
            reverse("geno:sysadmin-overview"),
        ]
        if hasattr(settings, "SHARE_PLOT") and settings.SHARE_PLOT:
            paths += [
                "member/overview/plot/",
                "share/overview/plot/",
            ]
        if django_apps.is_installed("portal"):
            paths += [
                "maintenance/check_portal_users/",
            ]
        for path in paths:
            if not path.startswith("/geno/"):
                path = f"/geno/{path}"
            response = self.client.get(path)
            if response.status_code != 200:
                print(f"FAILED PATH: {path} [{response.status_code}]")
            self.assertEqual(response.status_code, 200)

    def test_resident_unit_list_view(self):
        """Request the rental units (resident unit) list as XLSX and check response.

        Additionally verify that the calculated 'Bruttomiete' (ru_rent_total)
        takes contract.rent_reduction and contract.rent_reservation into
        account: ru_rent_total = rent_total - rent_reduction - rent_reservation.
        """
        # Ensure we have test data
        self.assertTrue(hasattr(self, "rentalunits"))
        self.assertTrue(len(self.rentalunits) > 0)

        # Attach a contract reduction to the first contract/rentalunit to test
        ru = self.rentalunits[0]
        contract = self.contracts[0]
        # Set monthly reductions (Fr./Monat)
        contract.rent_reduction = Decimal("50.00")
        contract.rent_reservation = Decimal("20.00")
        contract.save()

        expected_adjusted = ru.rent_total - Decimal("50.00") - Decimal("20.00")

        logged_in = self.client.login(username="superuser", password="secret")
        self.assertTrue(logged_in)
        path = reverse("geno:resident-list-units")
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)

        # export_data_to_xls returns an XLSX HttpResponse
        self.assertEqual(
            response.headers["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # Load workbook from response bytes and find the row for our rental unit
        wb = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
        ws = wb.active

        # Map header names to column indices
        header_map = {}
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            header_map[cell.value] = col

        # Expect headers 'Bezeichnung' and 'Bruttomiete' to be present
        self.assertIn("Bezeichnung", header_map)
        self.assertIn("Bruttomiete", header_map)

        bezeichnung_col = header_map["Bezeichnung"]
        brutto_col = header_map["Bruttomiete"]

        found = False
        for row in range(2, ws.max_row + 1):
            bezeichnung = ws.cell(row=row, column=bezeichnung_col).value
            if bezeichnung and str(ru.name) in str(bezeichnung):
                brutto_val = ws.cell(row=row, column=brutto_col).value
                # Convert to Decimal for comparison
                brutto_dec = Decimal(str(brutto_val)) if brutto_val is not None else None
                self.assertIsNotNone(brutto_dec)
                self.assertEqual(brutto_dec, expected_adjusted)
                found = True
                break

        self.assertTrue(found, f"Row for rental unit {ru} not found in XLSX export")

    @patch("geno.views.consolidate_invoices")
    def test_debtor_view_consolidate_invoices(self, mock_consolidate_invoices):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:debtor-list")
        response = self.client.post(path, data={"consolidate": ""})
        self.assertEqual(response.status_code, 200)
        mock_consolidate_invoices.assert_called_once()

        # Don't consolidate on search/filter
        response = self.client.post(path, data={"search": ""})
        self.assertEqual(response.status_code, 200)
        mock_consolidate_invoices.assert_called_once()

    @patch("geno.views.consolidate_invoices")
    def test_debtor_view_consolidate_invoices_detail(self, mock_consolidate_invoices):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:debtor-detail", args=["p", self.addresses[0].pk])
        response = self.client.post(path, data={"consolidate": ""})
        self.assertEqual(response.status_code, 200)
        mock_consolidate_invoices.assert_called_once_with(self.addresses[0])

        # Don't consolidate on search/filter
        response = self.client.post(path, data={"search": ""})
        self.assertEqual(response.status_code, 200)
        mock_consolidate_invoices.assert_called_once()

    def get_invoice_manual_data(self, preview=True, send_email=False):
        data = {
            "date": datetime.date.today().strftime("%d.%m.%Y"),
            "category": self.invoicecategories[0].pk,
            "address": self.addresses[0].pk,
            "extra_text": "Test extra_text",
            "form-0-date": datetime.date.today().strftime("%d.%m.%Y"),
            "form-0-text": "Test invoice item 1",
            "form-0-amount": "99.99",
            "form-TOTAL_FORMS": "1",
            "form-INITIAL_FORMS": "3",
        }
        if preview:
            data["submit_action_preview"] = ""
        else:
            data["submit_action_save"] = ""
        if send_email:
            data["send_email"] = "on"
        return data

    @patch("geno.views.create_qrbill", side_effect=create_fake_pdf, return_value=([], 0, None))
    def test_invoice_manual_process_preview(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        response = self.client.post(path, data=self.get_invoice_manual_data())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Content-Type"], "application/pdf")
        self.assertEqual(b"".join(response.streaming_content), b"Dummy PDF")
        self.assertEqual(Invoice.objects.count(), 0)

    @patch("geno.views.create_qrbill", side_effect=create_fake_pdf, return_value=([], 0, None))
    def test_invoice_manual_process_download(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        response = self.client.post(path, data=self.get_invoice_manual_data(preview=False))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Content-Type"], "application/pdf")
        self.assertEqual(b"".join(response.streaming_content), b"Dummy PDF")
        self.assertEqual(Invoice.objects.count(), 1)

    @patch(
        "geno.views.create_qrbill",
        side_effect=create_fake_pdf,
        return_value=([], 1, "test@example.test"),
    )
    def test_invoice_manual_process_send_email(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        response = self.client.post(
            path, data=self.get_invoice_manual_data(preview=False, send_email=True)
        )
        self.assertEqual(response.status_code, 200)
        self.assertInHTMLResponse(
            " mit QR-Rechnung an test@example.test geschickt. Rechnung_Member Invoice",
            response,
            raw=True,
        )
        self.assertEqual(Invoice.objects.count(), 1)

    @patch(
        "geno.views.create_qrbill",
        side_effect=create_fake_pdf,
        return_value=([], 0, "test@example.test"),
    )
    def test_invoice_manual_process_email_error(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        response = self.client.post(
            path, data=self.get_invoice_manual_data(preview=False, send_email=True)
        )
        self.assertEqual(response.status_code, 200)
        self.assertInHTMLResponse(
            "FEHLER beim Versenden des Emails ",
            response,
            raw=True,
        )
        self.assertEqual(Invoice.objects.count(), 0)

    @patch(
        "geno.views.create_qrbill",
        return_value=(["Error: TEST"], 1, "test@example.test"),
    )
    def test_invoice_manual_process_creation_error(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        response = self.client.post(
            path, data=self.get_invoice_manual_data(preview=False, send_email=False)
        )
        self.assertEqual(response.status_code, 200)
        self.assertInHTMLResponse(
            "Fehler beim Erzeugen der Rechnung für ",
            response,
            raw=True,
        )
        self.assertEqual(Invoice.objects.count(), 0)

    @patch("geno.views.create_qrbill", side_effect=create_fake_pdf, return_value=([], 0, None))
    def test_invoice_manual_address_and_contract_validation(self, _mock_create_qrbill):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        data = self.get_invoice_manual_data(preview=False, send_email=False)
        data["address"] = ""
        data["contract"] = self.contracts[0].pk
        response = self.client.post(path, data=data)
        self.assertInHTMLResponse(
            "Für diesen Rechnungstyp muss eine Adresse angegeben werden",
            response,
            raw=True,
        )
        data["category"] = InvoiceCategory.objects.get(name="Contract manual").pk
        data["address"] = self.addresses[0].pk
        data["contract"] = ""
        response = self.client.post(path, data=data)
        self.assertInHTMLResponse(
            "Für diesen Rechnungstyp muss ein Vertrag angegeben werden",
            response,
            raw=True,
        )
        data["contract"] = self.contracts[0].pk
        response = self.client.post(path, data=data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Content-Type"], "application/pdf")
        self.assertEqual(b"".join(response.streaming_content), b"Dummy PDF")
        self.assertEqual(Invoice.objects.count(), 1)
        invoice = Invoice.objects.first()
        self.assertEqual(invoice.contract, self.contracts[0])
        self.assertEqual(invoice.person, None)

    def test_invoice_manual_address_and_contract_form_field(self):
        self.client.login(username="superuser", password="secret")
        path = reverse("geno:invoice-manual")
        data = {"category": self.invoicecategories[0].pk}
        response = self.client.get(path, data=data)
        self.assertInHTMLResponse('<select name="address"', response, raw=True)
        self.assertNotInHTMLResponse('<select name="contract"', response, raw=True)
        data = {"category": InvoiceCategory.objects.get(name="Contract manual").pk}
        response = self.client.get(path, data=data)
        self.assertInHTMLResponse('<select name="contract"', response, raw=True)
        self.assertNotInHTMLResponse('<select name="address"', response, raw=True)


class InvoiceBatchGenerateViewTest(GenoAdminTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    @patch("geno.views.create_invoices")
    def test_invoice_batch_generate_view_item_count(self, mock_create_qrbill: Mock):
        view = InvoiceBatchGenerateView()
        view.kwargs = {"action": "create"}
        view.request = self.factory.post(reverse("geno:invoice-batch-generate"))
        view.request.user = self.su

        mock_create_qrbill.return_value = [{"info": "1 Rechnung für 1 Vertrag"}]
        view.result = view.process_action(dry_run=True)
        context = view.get_context_data()
        self.assertEqual(context["item_count"], 1)

        mock_create_qrbill.return_value = [{"info": "99 Rechnungen für 1 Vertrag"}]
        view.result = view.process_action(dry_run=True)
        context = view.get_context_data()
        self.assertEqual(context["item_count"], 99)

        mock_create_qrbill.return_value = [{"info": "10 Rechnungen für 9 Verträge"}]
        view.result = view.process_action(dry_run=True)
        context = view.get_context_data()
        self.assertEqual(context["item_count"], 10)


class Odt2PdfViewTest(GenoAdminTestCase):
    @tag("slow-test")
    def test_odt2pdf_view_singlefile(self):
        self.client.login(username="superuser", password="secret")
        with open("geno/tests/template_test.odt", "rb") as dummy_file:
            response = self.client.post(reverse("geno:odt2pdf"), {"file": dummy_file})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(isinstance(response, FileResponse))
        self.assertEqual(response.headers["Content-Type"], "application/pdf")
        self.assertInPDF(response.getvalue(), "Template-Test")

    @tag("slow-test")
    def test_odt2pdf_view_multifile(self):
        self.client.login(username="superuser", password="secret")
        with open("/tmp/test.zip", "wb") as zipfile_object:
            with zipfile.ZipFile(zipfile_object, "w") as archive:
                archive.write("geno/tests/template_test.odt", "testA.odt")
                archive.write("geno/tests/template_test.odt", "testB.odt")
        with open("/tmp/test.zip", "rb") as inputfile:
            response = self.client.post(reverse("geno:odt2pdf"), {"file": inputfile})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(isinstance(response, FileResponse))
        self.assertEqual(response.headers["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.getvalue())) as zip_file:
            self.assertInPDF(zip_file.read("testA.pdf"), "Template-Test")
            self.assertInPDF(zip_file.read("testB.pdf"), "Template-Test")


class ShareStatementViewTest(GenoAdminTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        stype_as = ShareType.objects.get(name="Anteilschein")
        Share.objects.all().delete()
        Share.objects.create(
            value=200,
            quantity=1,
            share_type=stype_as,
            payment_date=datetime.date(2020, 1, 1),
            name=cls.addresses[0],
        )
        Share.objects.create(
            value=200,
            quantity=5,
            share_type=stype_as,
            payment_date=datetime.date(2020, 1, 1),
            name=cls.addresses[1],
        )
        Share.objects.create(
            value=200,
            quantity=6,
            share_type=stype_as,
            payment_date=datetime.date(2020, 1, 1),
            name=cls.addresses[2],
        )
        Share.objects.create(
            value=200,
            quantity=10,
            share_type=stype_as,
            payment_date=datetime.date(2020, 1, 1),
            name=cls.addresses[3],
        )

    def test_get_object_skips(self):
        view = ShareStatementView()
        view.enddate = datetime.date(2020, 12, 31)
        obj = view.get_objects()
        # Default is no cutoff => no skips
        self.assertEqual(len(obj), 4)
        self.assertIn("Anzahl ignoriert=0", view.extra_description_info)

        with self.settings(GENO_SMALL_NUMBER_OF_SHARES_CUTOFF=5):
            obj = view.get_objects()
            self.assertEqual(len(obj), 2)
            self.assertIn("Anzahl ignoriert=2", view.extra_description_info)

            Share.objects.create(
                value=500,
                quantity=1,
                share_type=ShareType.objects.get(name="Depositenkasse"),
                payment_date=datetime.date(2020, 1, 1),
                name=self.addresses[1],
            )
            obj = view.get_objects()
            self.assertEqual(len(obj), 3)
            self.assertIn("Anzahl ignoriert=1", view.extra_description_info)

            view.address_id = self.addresses[0].pk
            obj = view.get_objects()
            self.assertEqual(len(obj), 1)
            self.assertIn("Anzahl ignoriert=0", view.extra_description_info)


class DocumentViewTest(GenoAdminTestCase):
    """Tests for DocumentCreateView, DocumentDownloadView and DocumentRecreateView."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # DocumentType and templates required by the document action views
        cls.shareconfirm_doctype = DocumentType.objects.create(
            name="shareconfirm", description="Share Confirmation"
        )
        cls.shareconfirm_doctype.templates.add(cls.contenttemplates[0])

        # A document with an existing file (fast-path download)
        cls.doc_with_file = Document.objects.create(
            name="Test_Doc_With_File.odt",
            doctype=cls.shareconfirm_doctype,
            template=cls.contenttemplates[0],
            data='{"test": "data"}',
            content_object=cls.shares[0],
        )
        cls.doc_with_file.file.save(
            "Test_Doc_With_File.odt", ContentFile(b"PDF existing content")
        )

        # A legacy document without a file (forces template re-render on download)
        cls.legacy_doc = Document.objects.create(
            name="Test_Legacy_Doc.odt",
            doctype=cls.shareconfirm_doctype,
            template=cls.contenttemplates[0],
            data='{"anrede": "Liebe Tests"}',
            content_object=cls.shares[0],
        )

    def setUp(self):
        super().setUp()
        # TODO: remove one of these patch paths, once `fill_pod_template()` has a final home
        patcher = patch("geno.views.fill_pod_template")
        self.mock_fill_template = patcher.start()
        self.mock_fill_template.side_effect = fill_template_effect
        self.addCleanup(patcher.stop)

        documents_patcher = patch("geno.documents.fill_pod_template")
        self.mock_fill_template_documents = documents_patcher.start()
        self.mock_fill_template_documents.side_effect = fill_template_effect
        self.addCleanup(documents_patcher.stop)

    def test_document_create_view(self):
        from django.contrib.contenttypes.models import ContentType

        url = reverse("geno:document-create", args=["shareconfirm", self.shares[0].pk])
        response = self.client.get(url, {"template": self.contenttemplates[0].pk})
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        self.assertEqual(content, b"odt mock\n")
        # A new Document should have been persisted for the share
        share_ct = ContentType.objects.get_for_model(self.shares[0])
        self.assertTrue(
            Document.objects.filter(
                doctype=self.shareconfirm_doctype,
                content_type=share_ct,
                object_id=self.shares[0].pk,
            ).exists()
        )

    def test_document_create_view_requires_template(self):
        url = reverse("geno:document-create", args=["shareconfirm", self.shares[0].pk])
        with self.assertRaises(RuntimeError):
            self.client.get(url)

    def test_document_create_view_permission_denied(self):
        self.client.logout()
        user = User.objects.create_user(username="no_perm", password="secret")
        self.client.login(username="no_perm", password="secret")
        url = reverse("geno:document-create", args=["shareconfirm", self.shares[0].pk])
        response = self.client.get(url, {"template": self.contenttemplates[0].pk})
        self.assertContains(response, "Keine Berechtigung", status_code=200)

    def test_document_download_view_with_file(self):
        url = reverse("geno:document-download", args=["shareconfirm", self.doc_with_file.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        self.assertEqual(content, b"PDF existing content")
        self.assertEqual(
            response.headers.get("Content-Disposition"),
            'attachment; filename="Test_Doc_With_File.odt"',
        )

    def test_document_download_view_without_file(self):
        url = reverse("geno:document-download", args=["shareconfirm", self.legacy_doc.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        self.assertEqual(content, b"odt mock\n")
        # The legacy Document itself should stay file-less
        self.legacy_doc.refresh_from_db()
        self.assertFalse(self.legacy_doc.file)

    def test_document_download_view_permission_denied(self):
        self.client.logout()
        user = User.objects.create_user(username="no_perm", password="secret")
        self.client.login(username="no_perm", password="secret")
        url = reverse("geno:document-download", args=["shareconfirm", self.doc_with_file.pk])
        response = self.client.get(url)
        self.assertContains(response, "Keine Berechtigung", status_code=200)

    def test_document_recreate_view(self):
        original_count = Document.objects.count()
        url = reverse("geno:document-recreate", args=["shareconfirm", self.legacy_doc.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        self.assertEqual(content, b"odt mock\n")
        # A new Document should have been created
        self.assertEqual(Document.objects.count(), original_count + 1)
        new_doc = Document.objects.exclude(pk=self.legacy_doc.pk).latest("pk")
        self.assertTrue(new_doc.file)
        # Original document should be untouched
        self.legacy_doc.refresh_from_db()
        self.assertEqual(self.legacy_doc.data, '{"anrede": "Liebe Tests"}')
        self.assertFalse(self.legacy_doc.file)

    def test_document_recreate_view_permission_denied(self):
        self.client.logout()
        user = User.objects.create_user(username="no_perm", password="secret")
        self.client.login(username="no_perm", password="secret")
        url = reverse("geno:document-recreate", args=["shareconfirm", self.legacy_doc.pk])
        response = self.client.get(url)
        self.assertContains(response, "Keine Berechtigung", status_code=200)
